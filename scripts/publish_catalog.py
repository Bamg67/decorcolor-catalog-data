# -*- coding: utf-8 -*-
"""Validate scan output and atomically publish JSON, CSV, XLSX, and manifest."""
from __future__ import annotations

import csv
import hashlib
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_RAW = "https://raw.githubusercontent.com/Bamg67/decorcolor-catalog-data/main/data/"
HEADERS = ["Товар", "Карточка RU", "Карточка UK", "Фото", "Цветов"]
for number in range(1, 4):
    HEADERS.extend([
        f"HEX {number}", f"Hue {number}", f"Hue от {number}", f"Hue до {number}",
        f"Доля {number}", f"L* от {number}", f"L* до {number}",
        f"DeltaE вероятное {number}", f"DeltaE проверка {number}",
    ])
HEADERS.append("Статус")


def validate(rows):
    if not rows:
        raise ValueError("catalog is empty")
    urls = [row.get("product_url") for row in rows]
    if any(not url for url in urls):
        raise ValueError("a product URL is missing")
    if len(urls) != len(set(urls)):
        raise ValueError("duplicate product URLs found")
    bad = [
        row for row in rows
        if not (
            (row.get("status") == "ok" and row.get("colors"))
            or (str(row.get("status", "")).startswith("review:") and not row.get("colors"))
        )
    ]
    if bad:
        raise ValueError("%d rows have an invalid color result" % len(bad))
    if any(len(row["colors"]) > 3 for row in rows):
        raise ValueError("a product has more than three colors")


def flat_row(row):
    values = [row["title"], row["product_url"], row.get("product_url_uk") or "",
              row["image_url"], len(row["colors"])]
    for index in range(3):
        color = row["colors"][index] if index < len(row["colors"]) else None
        if not color:
            values.extend([""] * 9)
            continue
        values.extend([
            color["hex"], color["hue_deg"], color["hue_from_deg"], color["hue_to_deg"],
            color["share"], color["lightness_from"], color["lightness_to"],
            color["delta_e_likely"], color["delta_e_review"],
        ])
    values.append("Готово" if row.get("status") == "ok" else "Требует проверки")
    return values


def write_atomic(path: Path, data: bytes):
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_bytes(data)
    os.replace(temporary, path)


def make_xlsx(rows, target: Path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Цвета товаров"
    sheet.sheet_view.showGridLines = False
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(flat_row(row))

    header_fill = PatternFill("solid", fgColor="7A273B")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = 52
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 15
    for column in range(5, len(HEADERS) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 12

    for row_index, row in enumerate(rows, 2):
        for column, label in ((2, "Открыть RU"), (3, "Открыть UK"), (4, "Открыть фото")):
            cell = sheet.cell(row_index, column)
            if cell.value:
                cell.hyperlink = cell.value
                cell.value = label
                cell.style = "Hyperlink"
        for color_index, column in enumerate((6, 15, 24)):
            if color_index >= len(row["colors"]):
                continue
            color = row["colors"][color_index]["hex"].lstrip("#")
            sheet.cell(row_index, column).fill = PatternFill("solid", fgColor=color)
            red, green, blue = (int(color[index:index + 2], 16) for index in (0, 2, 4))
            contrast = "000000" if 0.299 * red + 0.587 * green + 0.114 * blue > 150 else "FFFFFF"
            sheet.cell(row_index, column).font = Font(name="Arial", size=10, bold=True,
                                                      color=contrast)

    table = Table(displayName="ProductColors",
                  ref=f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    sheet.add_table(table)
    temporary = target.with_suffix(".xlsx.tmp")
    workbook.save(temporary)
    os.replace(temporary, target)


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: publish_catalog.py INPUT_JSON OUTPUT_DIR")
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    output.mkdir(parents=True, exist_ok=True)
    rows = json.loads(source.read_text(encoding="utf-8"))
    rows.sort(key=lambda row: (row["title"].casefold(), row["product_url"]))
    validate(rows)

    json_bytes = (json.dumps(rows, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    write_atomic(output / "catalog-colors.json", json_bytes)

    csv_path = output / "catalog-colors.csv"
    temporary_csv = csv_path.with_suffix(".csv.tmp")
    with temporary_csv.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, delimiter=";")
        writer.writerow(HEADERS)
        writer.writerows(flat_row(row) for row in rows)
    os.replace(temporary_csv, csv_path)
    make_xlsx(rows, output / "decor_opt_product_colors.xlsx")

    digest = hashlib.sha256(json_bytes).hexdigest()
    now = datetime.now(timezone.utc)
    manifest = {
        "schemaVersion": 1,
        "version": "%s.%s" % (now.strftime("%Y.%m.%d"), digest[:8]),
        "generatedAt": now.isoformat().replace("+00:00", "Z"),
        "products": len(rows),
        "coloredProducts": sum(bool(row.get("colors")) for row in rows),
        "reviewProducts": sum(not row.get("colors") for row in rows),
        "ukLinkedProducts": sum(bool(row.get("product_url_uk")) for row in rows),
        "catalogUrl": BASE_RAW + "catalog-colors.json",
        "csvUrl": BASE_RAW + "catalog-colors.csv",
        "sha256": digest,
        "deltaELikely": 10,
        "deltaEReview": 15,
    }
    write_atomic(output / "manifest.json", (json.dumps(manifest, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
    print(json.dumps(manifest, ensure_ascii=False))


if __name__ == "__main__":
    main()
